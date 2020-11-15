"""
@from：https://pythoneers.cn
@author: qq3330447288
@contact: erics1996@yeah.net
@software: PyCharm
@file: 在单元格中添加内容.py
@time: 2020/11/15 下午2:57
"""
from docx import Document

document = Document()
table = document.add_table(rows=2, cols=2)
cell = table.cell(0, 1)
cell.text = 'parrot, possibly dead'
document.save('./result/在单元格中添加内容.docx')
