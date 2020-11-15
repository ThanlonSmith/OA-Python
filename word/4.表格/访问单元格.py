"""
@from：https://pythoneers.cn
@author: qq3330447288
@contact: erics1996@yeah.net
@software: PyCharm
@file: 访问单元格.py
@time: 2020/11/15 下午2:51
"""
from docx import Document

document = Document()
table = document.add_table(rows=2, cols=2)
cell = table.cell(0, 1)
print(cell)  # <docx.table._Cell object at 0x7f01882b54c0>
